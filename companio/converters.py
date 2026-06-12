"""Office document → Markdown conversion layer.

Downloads are binary blobs that Claude Code's Read tool cannot parse.
This module converts xlsx/xlsm/docx/pptx to Markdown text files that
Read can handle. Unsupported or unconvertible files are returned as-is.

Dependencies are optional — if openpyxl / python-docx / python-pptx are
not installed, the converter gracefully returns the original path.
"""

from __future__ import annotations

import dataclasses
from pathlib import Path

from loguru import logger


_MAX_FILE_SIZE = 20 * 1024 * 1024  # 20 MB
_MAX_ROWS = 10_000
_MAX_CHARS = 200_000
_CONVERTIBLE_EXTS = frozenset({".xlsx", ".xlsm", ".docx", ".pptx"})


@dataclasses.dataclass
class ConversionResult:
    path: str
    original_path: str
    converted: bool
    meta: str  # human-readable summary, e.g. "3 sheets, 1240 rows"


def convert_if_needed(file_path: Path) -> ConversionResult:
    """Check extension and convert if supported. Returns original on failure."""
    original = str(file_path)
    ext = file_path.suffix.lower()

    if ext not in _CONVERTIBLE_EXTS:
        return ConversionResult(path=original, original_path=original, converted=False, meta="")

    if file_path.stat().st_size > _MAX_FILE_SIZE:
        logger.info("Skipping conversion for oversized file: {} ({:.1f} MB)", file_path.name, file_path.stat().st_size / 1024 / 1024)
        return ConversionResult(path=original, original_path=original, converted=False, meta="too large")

    try:
        if ext in (".xlsx", ".xlsm"):
            return _convert_excel(file_path)
        elif ext == ".docx":
            return _convert_word(file_path)
        elif ext == ".pptx":
            return _convert_ppt(file_path)
    except Exception as e:
        logger.warning("File conversion failed for {}: {}", file_path.name, e)

    return ConversionResult(path=original, original_path=original, converted=False, meta="conversion failed")


# ---------------------------------------------------------------------------
# Excel (.xlsx / .xlsm) → Markdown
# ---------------------------------------------------------------------------

def _read_excel_sheets(wb, max_rows: int) -> tuple[list[str], int, bool]:
    """Read all sheets from a workbook and return (lines, total_rows, truncated)."""
    lines: list[str] = []
    total_rows = 0
    truncated = False

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        lines.append(f"\n## Sheet: {sheet_name}\n")

        rows_data: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            if total_rows >= max_rows:
                truncated = True
                break
            cells = [_cell_to_str(c) for c in row]
            if any(cells):
                rows_data.append(cells)
                total_rows += 1

        if rows_data:
            lines.append(_rows_to_md_table(rows_data))

        if truncated:
            remaining_sheets = wb.sheetnames[wb.sheetnames.index(sheet_name) + 1:]
            if remaining_sheets:
                lines.append(f"\n... (행 제한 {max_rows}행 도달, 이후 시트 생략: {', '.join(remaining_sheets)})\n")
            else:
                lines.append(f"\n... ({max_rows}행 제한으로 잘림)\n")
            break

    return lines, total_rows, truncated


def _convert_excel(file_path: Path) -> ConversionResult:
    try:
        import openpyxl
    except ImportError:
        logger.warning("openpyxl not installed — xlsx/xlsm conversion unavailable. Install: pip install openpyxl")
        return ConversionResult(path=str(file_path), original_path=str(file_path), converted=False, meta="library missing")

    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    sheet_count = len(wb.sheetnames)
    sheet_lines, total_rows, truncated = _read_excel_sheets(wb, _MAX_ROWS)
    wb.close()

    # If read_only mode yielded suspiciously few rows, retry without it.
    # Some Excel files with complex merged cells or formatting break in read_only mode.
    if total_rows <= 2 and not truncated:
        logger.info("Excel read_only got only {} rows, retrying with read_only=False: {}", total_rows, file_path.name)
        wb = openpyxl.load_workbook(file_path, read_only=False, data_only=True)
        sheet_lines, total_rows, truncated = _read_excel_sheets(wb, _MAX_ROWS)
        wb.close()

    lines = [f"# {file_path.name}\n"] + sheet_lines
    content = "\n".join(lines)
    content = _truncate_chars(content)

    out_path = file_path.parent / f"{file_path.name}.md"
    out_path.write_text(content, encoding="utf-8")
    file_path.unlink(missing_ok=True)

    meta = f"{sheet_count} sheets, {total_rows} rows"
    if truncated:
        meta += " (truncated)"
    logger.info("Converted {} → {} ({})", file_path.name, out_path.name, meta)
    return ConversionResult(path=str(out_path), original_path=str(file_path), converted=True, meta=meta)


def _cell_to_str(value: object) -> str:
    if value is None:
        return ""
    return str(value).replace("|", "\\|").replace("\n", " ")


def _rows_to_md_table(rows: list[list[str]]) -> str:
    if not rows:
        return ""
    max_cols = max(len(r) for r in rows)
    for r in rows:
        while len(r) < max_cols:
            r.append("")

    # Skip meta/title rows that have fewer filled cells than the majority.
    # Heuristic: find the first row where >= 3 cells are non-empty (likely the real header).
    header_idx = 0
    if max_cols >= 3:
        for i, row in enumerate(rows):
            filled = sum(1 for c in row if c.strip())
            if filled >= min(3, max_cols):
                header_idx = i
                break

    prefix_lines: list[str] = []
    for row in rows[:header_idx]:
        text = " | ".join(c for c in row if c.strip())
        if text:
            prefix_lines.append(f"> {text}")

    data_rows = rows[header_idx:]
    if not data_rows:
        return "\n".join(prefix_lines) if prefix_lines else ""

    header = data_rows[0]
    separator = ["---"] * max_cols
    lines = prefix_lines + [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join(separator) + " |",
    ]
    for row in data_rows[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# Word (.docx) → Markdown
# ---------------------------------------------------------------------------

_HEADING_MAP = {
    "Heading 1": "# ",
    "Heading 2": "## ",
    "Heading 3": "### ",
    "Heading 4": "#### ",
}


def _convert_word(file_path: Path) -> ConversionResult:
    try:
        import docx as python_docx
    except ImportError:
        logger.warning("python-docx not installed — docx conversion unavailable. Install: pip install python-docx")
        return ConversionResult(path=str(file_path), original_path=str(file_path), converted=False, meta="library missing")

    doc = python_docx.Document(str(file_path))
    lines: list[str] = [f"# {file_path.name}\n"]
    para_count = 0

    para_by_el = {p._element: p for p in doc.paragraphs}
    tbl_by_el = {t._element: t for t in doc.tables}

    for element in doc.element.body:
        tag = element.tag.split("}")[-1]
        if tag == "p":
            para = para_by_el.get(element)
            if para is None:
                continue
            text = para.text.strip()
            if not text:
                lines.append("")
                continue
            prefix = _HEADING_MAP.get(para.style.name, "")
            lines.append(f"{prefix}{text}")
            para_count += 1
        elif tag == "tbl":
            tbl = tbl_by_el.get(element)
            if tbl is None:
                continue
            rows_data: list[list[str]] = []
            for row in tbl.rows:
                cells = [cell.text.strip().replace("|", "\\|").replace("\n", " ") for cell in row.cells]
                rows_data.append(cells)
            if rows_data:
                lines.append("")
                lines.append(_rows_to_md_table(rows_data))
                lines.append("")

    content = "\n".join(lines)
    content = _truncate_chars(content)

    out_path = file_path.parent / f"{file_path.name}.md"
    out_path.write_text(content, encoding="utf-8")
    file_path.unlink(missing_ok=True)

    meta = f"{para_count} paragraphs, {len(doc.tables)} tables"
    logger.info("Converted {} → {} ({})", file_path.name, out_path.name, meta)
    return ConversionResult(path=str(out_path), original_path=str(file_path), converted=True, meta=meta)


# ---------------------------------------------------------------------------
# PowerPoint (.pptx) → Markdown
# ---------------------------------------------------------------------------

def _convert_ppt(file_path: Path) -> ConversionResult:
    try:
        import pptx as python_pptx
        from pptx.enum.shapes import MSO_SHAPE_TYPE as _MSO_SHAPE_TYPE
    except ImportError:
        logger.warning("python-pptx not installed — pptx conversion unavailable. Install: pip install python-pptx")
        return ConversionResult(path=str(file_path), original_path=str(file_path), converted=False, meta="library missing")

    prs = python_pptx.Presentation(str(file_path))
    lines: list[str] = [f"# {file_path.name}\n"]
    slide_count = len(prs.slides)

    for i, slide in enumerate(prs.slides, 1):
        title = ""
        for shape in slide.shapes:
            if shape.has_text_frame and shape.shape_type is not None:
                if hasattr(shape, "placeholder_format") and shape.placeholder_format is not None:
                    if shape.placeholder_format.idx == 0:
                        title = shape.text_frame.text.strip()
                        break

        header = f"## Slide {i}" + (f": {title}" if title else "")
        lines.append(f"\n{header}\n")

        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text and text != title:
                        lines.append(f"- {text}")
            if shape.has_table:
                rows_data: list[list[str]] = []
                for row in shape.table.rows:
                    cells = [cell.text.strip().replace("|", "\\|").replace("\n", " ") for cell in row.cells]
                    rows_data.append(cells)
                if rows_data:
                    lines.append("")
                    lines.append(_rows_to_md_table(rows_data))
                    lines.append("")
            if hasattr(shape, "image") and _MSO_SHAPE_TYPE is not None and shape.shape_type == _MSO_SHAPE_TYPE.PICTURE:
                lines.append("- [이미지 생략]")
            if shape.has_chart:
                lines.append("- [차트 생략]")

        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"\n> 노트: {notes}")

    content = "\n".join(lines)
    content = _truncate_chars(content)

    out_path = file_path.parent / f"{file_path.name}.md"
    out_path.write_text(content, encoding="utf-8")
    file_path.unlink(missing_ok=True)

    meta = f"{slide_count} slides"
    logger.info("Converted {} → {} ({})", file_path.name, out_path.name, meta)
    return ConversionResult(path=str(out_path), original_path=str(file_path), converted=True, meta=meta)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _truncate_chars(content: str) -> str:
    if len(content) <= _MAX_CHARS:
        return content
    return content[:_MAX_CHARS] + f"\n\n... ({_MAX_CHARS}자 제한으로 잘림, 전체 {len(content)}자)"
